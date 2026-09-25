import os
import json
import random
from datetime import datetime, timedelta
import pandas as pd
from docx import Document
from faker import Faker

# ตั้งค่าจำนวนแถวขั้นต่ำ (10,000 รายการ)
NUM_RECORDS = 10000

# เริ่มใช้งาน Faker ภาษาไทย
fake = Faker('th_TH')

output_dir = "northwind_thai_large_data"
os.makedirs(output_dir, exist_ok=True)
print(f"🚀 กำลังสร้างชุดข้อมูลขนาดใหญ่สไตล์ไทยในโฟลเดอร์: {output_dir}\n")

# ---- [1. CSV Format] ข้อมูลลูกค้าชาวไทย 10,000 รายการ ----
print("⏳ กำลังสร้าง customers.csv...")
customer_ids = [f"CUST{str(i).zfill(5)}" for i in range(1, NUM_RECORDS + 1)]
customers_data = {
    "CustomerID": customer_ids,
    "CompanyName": [f"บริษัท {fake.company()} จำกัด" for _ in range(NUM_RECORDS)],
    "ContactName": [fake.name() for _ in range(NUM_RECORDS)],
    "City": [fake.city() for _ in range(NUM_RECORDS)],
    "Phone": [fake.phone_number() for _ in range(NUM_RECORDS)]
}
df_customers = pd.DataFrame(customers_data)
df_customers.to_csv(os.path.join(output_dir, "customers.csv"), index=False, encoding="utf-8-sig") # ใช้ utf-8-sig เพื่อให้ Excel อ่านภาษาไทยได้
print("✔ [CSV] สร้างไฟล์ 'customers.csv' เรียบร้อย (10,000 แถว)")

# ---- [2. JSON Format] ข้อมูลสินค้าประเภทต่างๆ 10,000 รายการ ----
print("⏳ กำลังสร้าง products.json...")
thai_products = ["ชาเขียวโออิชิ", "น้ำดื่มสิงห์", "มาม่าต้มยำกุ้ง", "กาแฟเบอร์ดี้", "ปลากระป๋องสามแม่ครัว", "สบู่ลักส์", "ผงซักฟอกบรีส", "น้ำยาล้างจานไลพอนเอฟ"]
products_data = []
for i in range(1, NUM_RECORDS + 1):
    base_prod = random.choice(thai_products)
    products_data.append({
        "ProductID": i,
        "ProductName": f"{base_prod} รสชาติพิเศษ {i}",
        "CategoryID": random.randint(1, 8),
        "Specs": {
            "QuantityPerUnit": f"แพ็ค {random.choice([4, 6, 12, 24])} ขวด/ซอง",
            "UnitPrice": round(random.uniform(10.0, 500.0), 2)
        }
    })
with open(os.path.join(output_dir, "products.json"), "w", encoding="utf-8") as f:
    json.dump(products_data, f, indent=4, ensure_ascii=False)
print("✔ [JSON] สร้างไฟล์ 'products.json' เรียบร้อย (10,000 แถว)")

# ---- [3. XLSX Format] ข้อมูลคำสั่งซื้อ 10,000 รายการ ----
print("⏳ กำลังสร้าง orders_2026.xlsx...")
start_date = datetime(2026, 1, 1)
orders_data = {
    "OrderID": [i for i in range(100001, 100001 + NUM_RECORDS)],
    "CustomerID": [random.choice(customer_ids) for _ in range(NUM_RECORDS)],
    "OrderDate": [(start_date + timedelta(days=random.randint(0, 260))).strftime("%Y-%m-%d") for _ in range(NUM_RECORDS)],
    "ShipperID": [random.randint(1, 3) for _ in range(NUM_RECORDS)],
    "Freight": [round(random.uniform(50.0, 1500.0), 2) for _ in range(NUM_RECORDS)]
}
df_orders = pd.DataFrame(orders_data)
df_orders.to_excel(os.path.join(output_dir, "orders_2026.xlsx"), index=False, engine='openpyxl')
print("✔ [XLSX] สร้างไฟล์ 'orders_2026.xlsx' เรียบร้อย (10,000 แถว)")

# ---- [4. TXT Format] รายละเอียดคำสั่งซื้อคั่นด้วย Pipe 10,000 รายการ ----
print("⏳ กำลังสร้าง order_details.txt...")
# สร้างหลายๆ แถวให้กับรายการสินค้าในบิล โดยผูกกับ OrderID ด้านบน
txt_lines = ["OrderID|ProductID|UnitPrice|Quantity|Discount"]
for i in range(1, NUM_RECORDS + 1):
    order_id = random.randint(100001, 100001 + NUM_RECORDS - 1)
    product_id = random.randint(1, NUM_RECORDS)
    unit_price = round(random.uniform(10.0, 500.0), 2)
    quantity = random.randint(1, 100)
    discount = random.choice([0.0, 0.05, 0.1, 0.15, 0.2])
    txt_lines.append(f"{order_id}|{product_id}|{unit_price}|{quantity}|{discount}")

with open(os.path.join(output_dir, "order_details.txt"), "w", encoding="utf-8") as f:
    f.write("\n".join(txt_lines))
print("✔ [TXT] สร้างไฟล์ 'order_details.txt' เรียบร้อย (10,000 แถว)")

# ---- [5. DOCX Format] นโยบายและรหัสผู้ขนส่ง (Shippers Reference) ----
print("⏳ กำลังสร้าง shippers_policy.docx...")
doc = Document()
doc.add_heading('Northwind Traders (Thailand) - คู่มือและการจัดการระวาง', 0)
doc.add_paragraph('เอกสารนี้จัดทำขึ้นในรูปข้อตกลงระดับบริหาร เพื่อกำหนดเกณฑ์การเชื่อมโยงข้อมูลโลจิสติกส์เข้าสู่คลังข้อมูลกลาง')

doc.add_heading('1. การจับคู่รหัสบริษัทขนส่งภายในประเทศ (Shipper ID)', level=1)
doc.add_paragraph('ในระเบียบจัดเก็บข้อมูลดิบ (Bronze Layer) ให้ยึดรหัส Mapping ตามที่ระบุในสัญญานี้เท่านั้น:')
doc.add_paragraph('• SHIP_ID: 1 -> ไปรษณีย์ไทย (Thailand Post)')
doc.add_paragraph('• SHIP_ID: 2 -> Kerry Express Thailand')
doc.add_paragraph('• SHIP_ID: 3 -> Flash Express')

doc.add_heading('2. ประกาศเพิ่มเติมเกี่ยวกับค่าธรรมเนียมและภาษี (Tax & Freight)', level=1)
doc.add_paragraph('เนื่องจากข้อมูลค่าระวาง (Freight) ในระบบสาขาต่างจังหวัดยังไม่ได้รวมภาษีมูลค่าเพิ่ม เมื่อระบบทรานส์ฟอร์มข้อมูลสู่ระดับ Gold Layer (Fact Table) เพื่อทำรายงานส่งผู้บริหารฝ่ายบัญชี ให้คำนวณปรับมูลค่าโดยทำการบวกเพิ่มอัตราภาษี 7% (คูณด้วย 1.07) เสมอ')

doc.save(os.path.join(output_dir, "shippers_policy.docx"))
print("✔ [DOCX] สร้างไฟล์ 'shippers_policy.docx' เรียบร้อย")

print(f"\n🎉 [เสร็จสิ้นการสร้างข้อมูลขนาดใหญ่] คุณได้ไฟล์ทดสอบปริมาณงานชุดภาษาไทยเรียบร้อยแล้วใน '{output_dir}/'")
